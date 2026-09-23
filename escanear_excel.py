"""
Extractor de estructura del Excel IEPPO.
Uso:
    python escanear_excel.py "ruta/al/archivo.xlsx"
Genera: ieppo_dump.txt
"""
import sys
import traceback
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def es_worksheet(ws) -> bool:
    return isinstance(ws, Worksheet)


def formatear_valor(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.replace("\n", "\\n")
        if len(v) > 300:
            v = v[:300] + "…"
        return v
    return v


def formatear_celda(c):
    if c.value is None:
        return None
    return formatear_valor(c.value)


def volcar_hoja(ws, out, max_filas=200, max_cols=80, titulo_extra=""):
    """Vuelca una worksheet con fórmulas y valores."""
    if not es_worksheet(ws):
        out.append(f"\n#### {ws.title} — NO es worksheet (tipo: {type(ws).__name__}), se omite")
        return

    try:
        dims = ws.dimensions
    except Exception:
        dims = "?"

    out.append(f"\n#### Hoja: {ws.title}  (dims: {dims}){titulo_extra}")
    out.append(f"  Estado: {ws.sheet_state}")

    max_r = min(ws.max_row or 0, max_filas)
    max_c = min(ws.max_column or 0, max_cols)

    for r in range(1, max_r + 1):
        fila_txt = []
        for c in range(1, max_c + 1):
            try:
                cell = ws.cell(row=r, column=c)
                v = formatear_celda(cell)
                if v is not None:
                    col = get_column_letter(c)
                    fila_txt.append(f"{col}{r}={v!r}")
            except Exception:
                pass
        if fila_txt:
            out.append("  " + " | ".join(fila_txt))


def listar_nombres_definidos(wb, out):
    out.append("\n## 2. NOMBRES DEFINIDOS")
    try:
        dn_container = wb.defined_names
        # En openpyxl >=3.1 es un dict-like
        if hasattr(dn_container, "items"):
            items = list(dn_container.items())
        else:
            items = [(dn.name, dn) for dn in dn_container.definedName]

        if not items:
            out.append("  (ninguno)")
            return

        for name, dn in items:
            try:
                valor = dn.value if hasattr(dn, "value") else str(dn)
            except Exception:
                valor = "?"
            out.append(f"  · {name} → {valor}")
    except Exception as e:
        out.append(f"  (error al leer nombres: {e})")
        out.append(traceback.format_exc())


def main(ruta_xlsx: str):
    ruta = Path(ruta_xlsx)
    if not ruta.exists():
        print(f"❌ No existe: {ruta}")
        sys.exit(1)

    print(f"📂 Leyendo: {ruta}")
    wb     = openpyxl.load_workbook(ruta, data_only=False)
    wb_val = openpyxl.load_workbook(ruta, data_only=True)

    out = []
    out.append("=" * 70)
    out.append("DUMP EXCEL IEPPO")
    out.append(f"Archivo: {ruta.name}")
    out.append("=" * 70)

    # ─────────────────────────────────────────────────────────
    # 1. HOJAS
    # ─────────────────────────────────────────────────────────
    out.append("\n## 1. HOJAS")
    for nombre in wb.sheetnames:
        try:
            ws = wb[nombre]
            tipo = type(ws).__name__
            estado = getattr(ws, "sheet_state", "?")
            dims = getattr(ws, "dimensions", "—")
            out.append(f"  · {nombre} [{tipo}] ({estado}) — {dims}")
        except Exception as e:
            out.append(f"  · {nombre} (error: {e})")

    # ─────────────────────────────────────────────────────────
    # 2. NOMBRES DEFINIDOS
    # ─────────────────────────────────────────────────────────
    listar_nombres_definidos(wb, out)

    # ─────────────────────────────────────────────────────────
    # 3. HOJAS CLAVE COMPLETAS
    # ─────────────────────────────────────────────────────────
    keywords = ["IEPPO", "INFORME", "CLAVE", "PUNT", "BARE",
                "RESP", "CALIF", "MATRIZ", "HOJA"]
    hojas_clave = [
        n for n in wb.sheetnames
        if any(k in n.upper() for k in keywords)
    ]

    out.append(f"\n## 3. HOJAS CLAVE DETECTADAS: {hojas_clave}")

    for nombre in hojas_clave:
        ws = wb[nombre]
        volcar_hoja(ws, out, max_filas=200, max_cols=100)

    # ─────────────────────────────────────────────────────────
    # 4. VALORES CALCULADOS (data_only=True)
    # ─────────────────────────────────────────────────────────
    out.append("\n## 4. VALORES CALCULADOS (hojas clave)")
    for nombre in hojas_clave:
        try:
            ws = wb_val[nombre]
            out.append(f"\n#### {nombre} (solo valores)")
            if not es_worksheet(ws):
                out.append("  (no es worksheet)")
                continue
            for r in range(1, min(ws.max_row or 0, 150) + 1):
                fila = []
                for c in range(1, min(ws.max_column or 0, 80) + 1):
                    try:
                        v = ws.cell(row=r, column=c).value
                        if v is not None:
                            col = get_column_letter(c)
                            fila.append(f"{col}{r}={formatear_valor(v)!r}")
                    except Exception:
                        pass
                if fila:
                    out.append("  " + " | ".join(fila))
        except Exception as e:
            out.append(f"  (error: {e})")

    # ─────────────────────────────────────────────────────────
    # 5. OTRAS HOJAS (encabezado)
    # ─────────────────────────────────────────────────────────
    out.append("\n## 5. OTRAS HOJAS (primeras 40 filas)")
    otras = [n for n in wb.sheetnames if n not in hojas_clave]
    for nombre in otras:
        ws = wb[nombre]
        volcar_hoja(ws, out, max_filas=40, max_cols=40)

    # ─────────────────────────────────────────────────────────
    # Escribir
    # ─────────────────────────────────────────────────────────
    salida = Path("ieppo_dump.txt")
    salida.write_text("\n".join(str(x) for x in out), encoding="utf-8")

    print(f"\n✅ Generado: {salida.resolve()}")
    print(f"   Tamaño:   {salida.stat().st_size / 1024:.1f} KB")
    print(f"   Hojas:    {len(wb.sheetnames)}")
    print(f"   Clave:    {hojas_clave}")
    print(f"\n👉 Pégame el contenido de ieppo_dump.txt (o súbelo como archivo).")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python escanear_excel.py <ruta_al_xlsx>")
        sys.exit(1)
    main(sys.argv[1])